import io, json, streamlit as st
from groq import Groq
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import db

MODEL = "openai/gpt-oss-120b"

def _client():
    key = st.secrets.get("GROQ_API_KEY", "")
    if not key:
        st.error("GROQ_API_KEY is missing. Add it in Streamlit: App settings -> Secrets.")
        st.stop()
    return Groq(api_key=key)

def llm(messages, as_json=False, temp=0.4):
    c = _client()
    base = dict(model=MODEL, messages=messages, temperature=temp, max_tokens=2048,
                extra_body={"reasoning_effort": "low"})
    if as_json:
        try:
            return c.chat.completions.create(
                response_format={"type": "json_object"}, **base).choices[0].message.content
        except Exception:
            pass  # fall back to plain mode and clean the output
    out = c.chat.completions.create(**base).choices[0].message.content
    return out.replace("```json", "").replace("```", "").strip() if as_json else out

# ---------- Ingestion Agent ----------
def extract(file):
    name, data = file.name, file.read()
    ext = name.lower().rsplit(".", 1)[-1]
    pages = []
    if ext == "pdf":
        for i, p in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
            pages.append((i, p.extract_text() or ""))
    elif ext == "pptx":
        from pptx import Presentation
        for i, s in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            pages.append((i, " ".join(sh.text for sh in s.shapes if sh.has_text_frame)))
    elif ext == "docx":
        from docx import Document
        pages.append((1, "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)))
    else:
        pages.append((1, data.decode("utf-8", "ignore")))
    return pages

def ingest(course_id, file):
    pages = extract(file)
    full = ""
    for pg, text in pages:
        full += text + "\n"
        for i in range(0, len(text), 900):
            piece = text[i:i + 1000].strip()
            if len(piece) > 40:
                db.run("INSERT INTO chunks(course_id,source,page,text) VALUES(?,?,?,?)",
                       (course_id, file.name, pg, piece))
    summary = llm([{"role": "user", "content":
        f"Summarize these lecture notes in 6 bullets and list 8 key terms:\n{full[:6000]}"}])
    db.run("INSERT INTO docs(course_id,name,summary) VALUES(?,?,?)", (course_id, file.name, summary))
    return len(pages)

# ---------- Retrieval Agent ----------
def retrieve(course_id, query, k=4):
    rows = db.q("SELECT * FROM chunks WHERE course_id=?", (course_id,))
    if not rows:
        return [], 0.0
    vec = TfidfVectorizer(stop_words="english").fit([r["text"] for r in rows] + [query])
    sims = cosine_similarity(vec.transform([query]), vec.transform([r["text"] for r in rows]))[0]
    top = sims.argsort()[::-1][:k]
    return [rows[i] for i in top], float(sims[top[0]])

# ---------- Tutor + Guardrail ----------
def tutor(course, question, mode, lang, history):
    chunks, score = retrieve(course["id"], question)
    if course["strict"] and score < 0.05:          # Guardrail: not grounded
        return "I couldn't find this in the course material. Please ask your teacher.", False, []
    ctx = "\n\n".join(f"[{c['source']} p.{c['page']}] {c['text']}" for c in chunks)
    style = {"Simple": "Explain simply with an example.",
             "Step-by-step": "Explain step by step.",
             "Socratic": "Guide with 2-3 leading questions; do not give the final answer immediately."}[mode]
    sys = (f"You are a university tutor. {style} Answer in {lang}. "
           "Use ONLY the context below and cite sources like [file p.N]. "
           "Do not complete graded assignments for the student.\n\nCONTEXT:\n" + ctx)
    msgs = [{"role": "system", "content": sys}] + history[-6:] + [{"role": "user", "content": question}]
    return llm(msgs, temp=0.5), True, chunks

# ---------- Quiz Agent ----------
def make_quiz(course_id, topic, n=5):
    chunks, _ = retrieve(course_id, topic, k=6)
    ctx = "\n".join(c["text"] for c in chunks)[:5000]
    p = (f"Create {n} multiple-choice questions on '{topic}' from this material:\n{ctx}\n"
         'Return JSON: {"questions":[{"q":"...","options":["A","B","C","D"],"answer":0,"explanation":"..."}]} '
         "where answer is the 0-based index of the correct option.")
    return json.loads(llm([{"role": "user", "content": p}], as_json=True))["questions"]

# ---------- Evaluator Agent ----------
def grade(questions, answers):
    correct = sum(1 for i, qn in enumerate(questions) if answers.get(i) == qn["answer"])
    return correct, len(questions)

# ---------- Orchestrator ----------
def route(text):
    r = llm([{"role": "user", "content":
        f"Classify into one word (question, quiz, summary): {text}"}], temp=0).strip().lower()
    return r if r in ("question", "quiz", "summary") else "question"
