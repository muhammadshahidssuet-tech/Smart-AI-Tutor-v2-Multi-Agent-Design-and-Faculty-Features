import io, json
import pandas as pd
import db

def build_report(course_id, quiz):
    """One row per enrolled student: name, enrollment, weak topic, marks, status bar, Q1..Qn right/wrong."""
    questions = json.loads(quiz["data"])
    studs = db.q("""SELECT u.id, COALESCE(u.full_name, u.username) AS name,
                           COALESCE(u.enrollment_no, '-') AS enr
                    FROM enrollments e JOIN users u ON u.id = e.user_id
                    WHERE e.course_id=? ORDER BY u.enrollment_no""", (course_id,))
    this = {a["user_id"]: a for a in db.q(
        "SELECT * FROM attempts WHERE course_id=? AND quiz_id=?", (course_id, quiz["id"]))}
    allatt = db.q("SELECT user_id, topic, correct, total FROM attempts WHERE course_id=?", (course_id,))
    rows = []
    for s in studs:
        per = {}
        for x in allatt:
            if x["user_id"] == s["id"]:
                v = per.setdefault(x["topic"], [0, 0]); v[0] += x["correct"]; v[1] += x["total"]
        low = {k: v[0] / v[1] for k, v in per.items() if v[1] and v[0] / v[1] < 0.6}
        weak = min(low, key=low.get) if low else ("None" if per else "-")
        row = {"Student Name": s["name"], "Enrollment No.": s["enr"], "Weak Topic": weak}
        a = this.get(s["id"])
        if a:
            pct = round(a["correct"] / a["total"] * 100) if a["total"] else 0
            row["Marks Obtained"] = f"{a['correct']}/{a['total']}"
            row["Percentage"] = pct
            row["Status Bar"] = "█" * (pct // 10) + "░" * (10 - pct // 10) + f" {pct}%"
            saved = json.loads(a["answers"]) if a["answers"] else None
            for i, qn in enumerate(questions):
                row[f"Q{i+1}"] = "N/A" if saved is None else (
                    "Correct" if saved.get(str(i)) == qn["answer"] else "Wrong")
        else:
            row["Marks Obtained"] = "Not attempted"
            row["Percentage"] = 0
            row["Status Bar"] = "░" * 10 + " 0%"
            for i in range(len(questions)):
                row[f"Q{i+1}"] = "-"
        rows.append(row)
    return pd.DataFrame(rows)

def to_csv(df):
    return df.to_csv(index=False).encode("utf-8-sig")   # utf-8-sig so Excel shows the bar symbols

def to_excel(df):
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Quiz Report")
        ws = w.sheets["Quiz Report"]
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor="4F46E5")
            c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center")
        green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.value == "Correct": c.fill = green
                elif c.value == "Wrong": c.fill = red
        if len(df):
            L = get_column_letter(list(df.columns).index("Percentage") + 1)
            ws.conditional_formatting.add(f"{L}2:{L}{len(df) + 1}", DataBarRule(
                start_type="num", start_value=0, end_type="num", end_value=100, color="4F46E5"))
        for i, col in enumerate(df.columns, 1):
            width = max([len(str(col))] + [len(str(v)) for v in df[col]]) + 3
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 32)
        ws.freeze_panes = "C2"
    return buf.getvalue()
